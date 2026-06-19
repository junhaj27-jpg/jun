import json
import tempfile
import unittest
from pathlib import Path

from docx import Document

from agents.document_merge.agent import DocumentMergeAgent
from tools.result import success_result


class FakeLLMClient:
    def chat(self, messages, **kwargs):
        return success_result(
            {
                "meeting_change_items": [
                    {
                        "change_type": "ADD",
                        "item": {"req_id": "REQ-NEW", "name": "추가 요구사항"},
                        "search_targets": "WEB",
                        "search_query": "표준 검색",
                    }
                ]
            }
        )


class RoutingLLMClient:
    def __init__(self) -> None:
        self.calls = []

    def chat(self, messages, **kwargs):
        self.calls.append(messages)
        system = messages[0]["content"]
        if "회의록" in system and "ADD, UPDATE, DELETE" in system:
            return success_result(
                {
                    "meeting_change_items": [
                        {
                            "change_type": "UPDATE",
                            "target_id": "REQ-001",
                            "content": {"req_id": "REQ-001", "name": "회의록 반영"},
                            "search_targets": "RAG",
                            "search_query": "정책 검색",
                        },
                        {
                            "change_type": "ADD",
                            "item": {"req_id": "REQ-NEW", "name": "신규"},
                        },
                    ]
                }
            )
        if "검색 결과" in system:
            return success_result({"content": {"req_id": "REQ-001", "name": "검색 반영"}})
        if "기존 item" in system:
            return success_result(
                {
                    "change_type": "UPDATE",
                    "item": {"req_id": "REQ-001", "name": "병렬 반영"},
                }
            )
        return success_result({})


class DocumentMergeAgentTest(unittest.TestCase):
    def test_srs_create_uses_rfp_parser_meeting_llm_and_search_router(self) -> None:
        with tempfile.TemporaryDirectory() as root:
            meeting = Path(root) / "meeting.txt"
            meeting.write_text("신규 요구사항을 추가한다.", encoding="utf-8")
            search_calls = []

            def rfp_parser(path):
                return success_result(
                    {"functional_requirements": [{"req_id": "REQ-001", "name": "기존"}]}
                )

            def search_tool(query, **kwargs):
                search_calls.append((query, kwargs))
                return success_result(
                    {"normalized_results": [{"source": "WEB", "content": "검색 결과"}]}
                )

            state = {
                "docs_cd": "SRS",
                "udt_yn": "N",
                "base_rfp_path": str(Path(root) / "rfp.pdf"),
                "input_file_paths": [str(meeting)],
                "agent_outputs": {},
            }
            result = DocumentMergeAgent(
                llm_client=FakeLLMClient(),
                rfp_parser=rfp_parser,
                search_tool=search_tool,
            ).execute(state)

            self.assertEqual(result["status"], "SUCCESS")
            self.assertEqual(len(result["integrated_requirement_json_list"]), 2)
            self.assertEqual(search_calls[0][1]["search_targets"], "WEB")
            self.assertIs(state["agent_outputs"]["document_merge_agent"], result)
            self.assertNotIn("integrated_requirement_json_list", state)

    def test_srs_create_uses_parallel_item_merge_and_embedding_writer(self) -> None:
        with tempfile.TemporaryDirectory() as root:
            meeting = Path(root) / "meeting.txt"
            meeting.write_text("REQ-001을 수정하고 신규 요구사항을 추가한다.", encoding="utf-8")
            llm = RoutingLLMClient()
            embedding_calls = []

            def rfp_parser(path):
                return success_result(
                    {
                        "functional_requirements": [
                            {
                                "req_id": "REQ-001",
                                "name": "기존",
                                "requirement_type": "기능",
                            },
                            {
                                "req_id": "NFR-001",
                                "name": "보안",
                                "requirement_type": "보안",
                            },
                        ]
                    }
                )

            def search_tool(query, **kwargs):
                return success_result(
                    {"normalized_results": [{"source_kind": "RAG", "content": "검색 결과"}]}
                )

            def embedding_writer(requirements, **kwargs):
                embedding_calls.append((requirements, kwargs))
                return success_result({"stored_count": 1})

            result = DocumentMergeAgent(
                llm_client=llm,
                rfp_parser=rfp_parser,
                search_tool=search_tool,
                embedding_writer=embedding_writer,
                max_parallel_workers=2,
            ).execute(
                {
                    "project_sn": 1,
                    "docs_cd": "SRS",
                    "udt_yn": "N",
                    "base_rfp_path": str(Path(root) / "rfp.docx"),
                    "input_file_paths": [str(meeting)],
                    "agent_outputs": {},
                }
            )

            self.assertEqual(result["status"], "SUCCESS")
            names = {item.get("req_id"): item.get("name") for item in result["integrated_requirement_json_list"]}
            self.assertEqual(names["REQ-001"], "병렬 반영")
            self.assertEqual(names["REQ-NEW"], "신규")
            self.assertGreaterEqual(len(llm.calls), 3)
            self.assertTrue(embedding_calls)
            embedded_ids = {item.get("req_id") for item in embedding_calls[0][0]}
            self.assertIn("NFR-001", embedded_ids)
            self.assertEqual(embedding_calls[0][1]["project_sn"], 1)

    def test_other_create_loads_requirements_and_reference_documents(self) -> None:
        with tempfile.TemporaryDirectory() as root:
            root_path = Path(root)
            requirements = root_path / "requirements.json"
            erd = root_path / "erd.json"
            interface = root_path / "interface.json"
            requirements.write_text(
                json.dumps({"requirement_json_list": [{"req_id": "REQ-001"}]}),
                encoding="utf-8",
            )
            erd.write_text(json.dumps({"tables": [{"table_id": "T1"}]}), encoding="utf-8")
            interface.write_text(
                json.dumps({"screens": [{"screen_id": "SCR-1"}]}), encoding="utf-8"
            )

            db = DocumentMergeAgent().execute(
                {
                    "docs_cd": "DB",
                    "udt_yn": "N",
                    "base_requirement_json_path": str(requirements),
                    "erd_file_path": str(erd),
                }
            )
            ts = DocumentMergeAgent().execute(
                {
                    "docs_cd": "TS",
                    "udt_yn": "N",
                    "base_requirement_json_path": str(requirements),
                    "interface_file_path": str(interface),
                }
            )

            self.assertEqual(db["reference_erd_json_list"], [{"table_id": "T1"}])
            self.assertEqual(ts["reference_interface_json_list"], [{"screen_id": "SCR-1"}])

    def test_db_create_loads_reference_erd_from_docx_entity_tables(self) -> None:
        with tempfile.TemporaryDirectory() as root:
            root_path = Path(root)
            requirements = root_path / "requirements.json"
            erd = root_path / "erd.docx"
            requirements.write_text(
                json.dumps({"requirement_json_list": [{"req_id": "REQ-001"}]}),
                encoding="utf-8",
            )
            document = Document()
            table = document.add_table(rows=4, cols=10)
            table.cell(0, 2).text = "ENT-001"
            table.cell(0, 7).text = "사용자"
            table.cell(1, 4).text = "사용자 정보를 관리하는 테이블입니다."
            table.rows[2].cells[0].text = "속성명"
            table.rows[2].cells[1].text = "동의어"
            table.rows[2].cells[2].text = "데이터타입"
            table.rows[3].cells[0].text = "USER_SN"
            table.rows[3].cells[1].text = "사용자 번호"
            table.rows[3].cells[2].text = "BIGINT"
            table.rows[3].cells[4].text = "Y"
            table.rows[3].cells[5].text = "Y"
            document.save(erd)

            result = DocumentMergeAgent().execute(
                {
                    "docs_cd": "DB",
                    "udt_yn": "N",
                    "base_requirement_json_path": str(requirements),
                    "erd_file_path": str(erd),
                }
            )

            self.assertEqual(result["status"], "SUCCESS")
            self.assertEqual(result["reference_erd_json_list"][0]["logical_name"], "사용자")
            self.assertEqual(result["reference_erd_json_list"][0]["columns"][0]["physical_name"], "USER_SN")

    def test_other_create_loads_requirement_json_from_final_document_wrapper(self) -> None:
        with tempfile.TemporaryDirectory() as root:
            requirements = Path(root) / "requirements.json"
            requirements.write_text(
                json.dumps(
                    {
                        "result": {
                            "final_document_json": {
                                "docs_cd": "SRS",
                                "requirement_json_list": [
                                    {
                                        "req_id": "REQ-001",
                                        "requirement_type": "기능",
                                        "detail_text": "데이터를 관리한다.",
                                    }
                                ],
                            }
                        }
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            result = DocumentMergeAgent().execute(
                {
                    "docs_cd": "ERD",
                    "udt_yn": "N",
                    "base_requirement_json_path": str(requirements),
                }
            )

            self.assertEqual(result["status"], "SUCCESS")
            self.assertEqual(result["integrated_requirement_json_list"][0]["req_id"], "REQ-001")

    def test_other_create_filters_requirements_by_docs_cd(self) -> None:
        with tempfile.TemporaryDirectory() as root:
            requirements = Path(root) / "requirements.json"
            requirements.write_text(
                json.dumps(
                    {
                        "requirement_json_list": [
                            {"req_id": "REQ-F", "requirement_type": "기능", "detail_text": "로그인"},
                            {"req_id": "REQ-I", "requirement_type": "인터페이스", "detail_text": "화면"},
                            {"req_id": "REQ-P", "requirement_type": "성능", "detail_text": "응답"},
                        ]
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            result = DocumentMergeAgent().execute(
                {
                    "docs_cd": "INTERFACE",
                    "udt_yn": "N",
                    "base_requirement_json_path": str(requirements),
                }
            )

            ids = [item["req_id"] for item in result["integrated_requirement_json_list"]]
            self.assertEqual(ids, ["REQ-I"])

    def test_update_modes_keep_or_merge_outputs_without_top_level_writes(self) -> None:
        with tempfile.TemporaryDirectory() as root:
            root_path = Path(root)
            existing = root_path / "existing.json"
            meeting = root_path / "meeting.txt"
            existing.write_text(
                json.dumps({"items": [{"id": "A", "name": "기존"}]}),
                encoding="utf-8",
            )
            meeting.write_text("기존 내용을 변경한다.", encoding="utf-8")

            structural_state = {
                "docs_cd": "ARCH",
                "udt_yn": "Y",
                "existing_output_path": str(existing),
                "input_file_paths": [str(meeting)],
            }
            document_state = {
                "docs_cd": "INTERFACE",
                "udt_yn": "Y",
                "existing_output_path": str(existing),
                "input_file_paths": [str(meeting)],
            }
            structural = DocumentMergeAgent().execute(structural_state)
            document = DocumentMergeAgent().execute(document_state)

            self.assertIn("existing_output_raw_json", structural)
            self.assertIn("meeting_change_items", structural)
            self.assertIn("integrated_artifact_json_list", document)
            self.assertNotIn("existing_output_raw_json", document)
            self.assertNotIn("integrated_artifact_json_list", document_state)

    def test_erd_update_docx_uses_erd_parser_for_existing_output(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as root:
            root_path = Path(root)
            existing = root_path / "existing_erd.docx"
            meeting = root_path / "meeting.txt"
            document = Document()
            table = document.add_table(rows=4, cols=10)
            table.cell(0, 2).text = "ENT-001"
            table.cell(0, 7).text = "사용자"
            table.cell(1, 4).text = "사용자 정보를 관리합니다."
            table.rows[2].cells[0].text = "속성명"
            table.rows[2].cells[1].text = "동의어"
            table.rows[2].cells[2].text = "타입"
            table.rows[3].cells[0].text = "user_id"
            table.rows[3].cells[1].text = "사용자 ID"
            table.rows[3].cells[2].text = "BIGINT"
            table.rows[3].cells[4].text = "Y"
            table.rows[3].cells[5].text = "PK"
            document.save(existing)
            meeting.write_text("사용자 엔티티 설명을 보완한다.", encoding="utf-8")

            result = DocumentMergeAgent().execute(
                {
                    "docs_cd": "ERD",
                    "udt_yn": "Y",
                    "existing_output_path": str(existing),
                    "input_file_paths": [str(meeting)],
                }
            )

            self.assertEqual(result["status"], "SUCCESS")
            self.assertEqual(result["existing_output_raw_json"]["tables"][0]["logical_name"], "사용자")

    def test_update_fallback_meeting_text_is_not_appended_as_artifact_item(self) -> None:
        with tempfile.TemporaryDirectory() as root:
            root_path = Path(root)
            existing = root_path / "existing.json"
            meeting = root_path / "meeting.txt"
            existing.write_text(
                json.dumps(
                    {
                        "requirement_json_list": [
                            {
                                "requirement_id": "SFR-001",
                                "requirement_name": "로그인",
                                "requirement_type": "기능",
                                "description": "로그인한다.",
                                "source": ["SFR-001"],
                                "validation_criteria": ["로그인 성공 여부 확인"],
                            }
                        ]
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            meeting.write_text("변경 요구사항 검토 회의록 원문", encoding="utf-8")

            result = DocumentMergeAgent().execute(
                {
                    "docs_cd": "SRS",
                    "udt_yn": "Y",
                    "existing_output_path": str(existing),
                    "input_file_paths": [str(meeting)],
                }
            )

            self.assertEqual(result["status"], "SUCCESS")
            self.assertEqual(len(result["integrated_artifact_json_list"]), 1)
            self.assertTrue(
                all(
                    isinstance(item, dict)
                    for item in result["integrated_artifact_json_list"]
                )
            )

    def test_missing_required_input_returns_failed_output_in_agent_outputs(self) -> None:
        state = {"docs_cd": "SRS", "udt_yn": "N"}
        result = DocumentMergeAgent().execute(state)

        self.assertEqual(result["status"], "FAILED")
        self.assertEqual(result["failure_type"], "SRS_RFP_MISSING")
        self.assertEqual(state["agent_outputs"]["document_merge_agent"], result)


if __name__ == "__main__":
    unittest.main()
